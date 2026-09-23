"""
Tests for Track C PowerPoint Deck Engine (services.api.deck_engine).
"""

from __future__ import annotations

import io
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from pptx import Presentation
from sqlmodel import Session, SQLModel, create_engine
from sqlmodel.pool import StaticPool

from services.api.deck_engine import deck_engine
from services.api.models import (
    DataFrameCheckpoint,
    DeckConfigRequest,
    Session as DbSession,
)


@pytest.fixture
def test_db():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    SQLModel.metadata.create_all(engine)
    with Session(engine) as session:
        yield session


@pytest.fixture
def mock_session_with_data(test_db: Session, tmp_path: Path):
    """Creates a test session with a versioned Parquet checkpoint."""
    session = DbSession(
        id="test-deck-session-123",
        title="Q3 Enterprise Cohort Performance",
        active_dataframe_version="df_v0",
        workspace_id="default_ws",
    )
    test_db.add(session)
    test_db.commit()

    # Create dummy DataFrame
    np.random.seed(42)
    df = pd.DataFrame({
        "revenue": np.random.uniform(100, 1000, 50),
        "subscribers": np.random.randint(10, 500, 50),
        "churn_rate": np.random.uniform(0.01, 0.15, 50),
        "cohort": np.random.choice(["Alpha", "Beta", "Gamma"], 50),
    })

    parquet_path = tmp_path / "df_v0.parquet"
    df.to_parquet(parquet_path, index=False)

    chk = DataFrameCheckpoint(
        checkpoint_id="chk_001",
        session_id="test-deck-session-123",
        version_tag="df_v0",
        parquet_storage_path=str(parquet_path),
        row_count=len(df),
        column_count=len(df.columns),
        operation_summary="Initial ingestion",
        is_active=True,
    )
    test_db.add(chk)
    test_db.commit()

    return session


def test_deck_preview_generation(test_db: Session, mock_session_with_data):
    """Test generating structured slide outlines for web studio preview."""
    config = DeckConfigRequest(theme="dark", title="Custom Deck Title")
    preview = deck_engine.generate_deck_preview(test_db, "test-deck-session-123", config)

    assert preview.session_id == "test-deck-session-123"
    assert preview.title == "Custom Deck Title"
    assert preview.total_slides >= 4
    assert len(preview.slides) == preview.total_slides

    slide_types = [s.slide_type for s in preview.slides]
    assert "title" in slide_types
    assert "hygiene" in slide_types
    assert "summary" in slide_types


def test_deck_pptx_compilation(test_db: Session, mock_session_with_data):
    """Test compiling valid .pptx binary with native charts and shapes."""
    for theme in ["dark", "light", "navy"]:
        config = DeckConfigRequest(theme=theme, title="Boardroom Quarterly Review")
        pptx_bytes = deck_engine.build_deck(test_db, "test-deck-session-123", config)

        assert isinstance(pptx_bytes, bytes)
        assert len(pptx_bytes) > 5000  # Non-trivial presentation size

        # Verify ZIP/OfficeOpenXML signature
        assert pptx_bytes[:4] == b"PK\x03\x04"

        # Verify presentation can be parsed back by python-pptx
        prs = Presentation(io.BytesIO(pptx_bytes))
        assert len(prs.slides) >= 4
